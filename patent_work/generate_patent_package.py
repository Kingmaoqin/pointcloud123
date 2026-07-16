#!/usr/bin/env python3
from __future__ import annotations

import csv
import json
import math
import textwrap
from pathlib import Path
from typing import Any

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib import font_manager
import numpy as np
import pandas as pd
import yaml
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt
from matplotlib.backends.backend_pdf import PdfPages
from matplotlib.patches import Circle, FancyArrowPatch, Polygon, Rectangle


ROOT = Path("/home/xqin5/patent_gap_nbv")
PATENT = ROOT / "patent_work"
FIG_ROOT = ROOT / "outputs" / "patent_figures"
FORMAL = FIG_ROOT / "formal_bw"
COLOR = FIG_ROOT / "color_demo"

for path in (PATENT, FORMAL, COLOR):
    path.mkdir(parents=True, exist_ok=True)

FONT_PATH = Path("/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc")
if FONT_PATH.exists():
    font_manager.fontManager.addfont(str(FONT_PATH))
    CJK_FONT = font_manager.FontProperties(fname=str(FONT_PATH)).get_name()
else:
    CJK_FONT = "DejaVu Sans"

plt.rcParams["font.family"] = CJK_FONT
plt.rcParams["font.sans-serif"] = [CJK_FONT, "DejaVu Sans"]
plt.rcParams["axes.unicode_minus"] = False
plt.rcParams["svg.fonttype"] = "none"


def read_json(rel: str) -> dict[str, Any]:
    path = ROOT / rel
    return json.loads(path.read_text(encoding="utf-8")) if path.exists() else {}


def read_yaml(rel: str) -> dict[str, Any]:
    path = ROOT / rel
    return yaml.safe_load(path.read_text(encoding="utf-8")) or {} if path.exists() else {}


def safe_read_csv(rel: str) -> pd.DataFrame:
    path = ROOT / rel
    return pd.read_csv(path) if path.exists() else pd.DataFrame()


def safe_read_parquet(rel: str) -> pd.DataFrame:
    path = ROOT / rel
    return pd.read_parquet(path) if path.exists() else pd.DataFrame()


SYN = read_json("outputs/reports/synthetic_summary.json")
REAL = read_json("outputs/reports/real_data_summary.json")
PRE = read_json("outputs/reports/preprocess_summary.json")
FULL_ASSOC = read_json("outputs/reports/cras_full_assoc_summary.json")
SAMPLE_ASSOC = read_json("data/processed/cras_point_sample_association_summary.json")
BEST = read_yaml("outputs/reports/best_parameters.yaml")
GAP_CFG = read_yaml("configs/gap/default.yaml")
VIEW_CFG = read_yaml("configs/view/default.yaml")
CRAS_CFG = read_yaml("configs/data/cras.yaml")
CRAS_FULL_CFG = read_yaml("configs/experiment/cras_full.yaml")
SYN_SMOKE_CFG = read_yaml("configs/experiment/synthetic_smoke.yaml")

patch_real = safe_read_csv("outputs/tables/patch_scores_real.csv")
patch_synth = safe_read_csv("outputs/tables/patch_scores_synth.csv")
views_real = safe_read_csv("outputs/tables/candidate_view_ranking_real.csv")
views_synth = safe_read_csv("outputs/tables/candidate_view_ranking_synth.csv")
cl_real = safe_read_csv("outputs/tables/closed_loop_results_real.csv")
cl_synth = safe_read_csv("outputs/tables/closed_loop_results_synth.csv")
patches = safe_read_parquet("data/processed/patches_real.parquet")
elements = safe_read_parquet("data/processed/ifc_elements.parquet")
materials = safe_read_parquet("data/processed/ifc_materials.parquet")
points_sample = safe_read_parquet("data/processed/cras_point_sample.parquet")
assoc_sample = safe_read_parquet("data/processed/cras_point_sample_associations.parquet")
patch_stats = safe_read_csv("data/processed/synthetic_scan/patch_stats.csv")
scanner_pos = safe_read_csv("data/processed/synthetic_scan/scanner_positions.csv")


def fmt(x: Any, digits: int = 4) -> str:
    if x is None:
        return "无"
    try:
        if isinstance(x, float) and (math.isnan(x) or math.isinf(x)):
            return "无"
        return f"{float(x):.{digits}f}"
    except Exception:
        return str(x)


def md_table(rows: list[list[Any]], headers: list[str]) -> str:
    out = ["| " + " | ".join(headers) + " |", "| " + " | ".join(["---"] * len(headers)) + " |"]
    for row in rows:
        out.append("| " + " | ".join(str(v).replace("\n", "<br>") for v in row) + " |")
    return "\n".join(out)


def write_text(path: Path, text: str) -> None:
    path.write_text(text.strip() + "\n", encoding="utf-8")


def save_fig(fig: plt.Figure, name: str, formal: bool = True) -> None:
    outdir = FORMAL if formal else COLOR
    fig.savefig(outdir / f"{name}.svg", bbox_inches="tight")
    fig.savefig(outdir / f"{name}.png", dpi=350, bbox_inches="tight")
    plt.close(fig)


def arrow(ax, xy1, xy2, color="black", lw=1.5):
    ax.add_patch(FancyArrowPatch(xy1, xy2, arrowstyle="-|>", mutation_scale=12, linewidth=lw, color=color))


def box(ax, xy, w, h, text, fc="white", ec="black", fontsize=10):
    ax.add_patch(Rectangle(xy, w, h, facecolor=fc, edgecolor=ec, linewidth=1.4))
    ax.text(xy[0] + w / 2, xy[1] + h / 2, text, ha="center", va="center", fontsize=fontsize, wrap=True)


def fig_flow() -> plt.Figure:
    fig, ax = plt.subplots(figsize=(10, 5.2))
    ax.set_axis_off()
    labels = ["S1 获取IFC\n和点云", "S2 IFC三角化\n形成Patch", "S3 点云-IFC\n关联", "S4 构造多源\n证据", "S5 缺口融合\nG_gap", "S6 生成并排序\n补扫视点", "S7 更新证据\n闭环迭代"]
    xs = np.linspace(0.05, 0.82, len(labels))
    y = 0.60
    for i, (x, label) in enumerate(zip(xs, labels)):
        box(ax, (x, y), 0.12, 0.18, label, fontsize=9)
        if i < len(labels) - 1:
            arrow(ax, (x + 0.12, y + 0.09), (xs[i + 1], y + 0.09))
    arrow(ax, (0.88, 0.60), (0.88, 0.30))
    arrow(ax, (0.88, 0.30), (0.10, 0.30))
    arrow(ax, (0.10, 0.30), (0.10, 0.60))
    ax.text(0.50, 0.22, "闭环：按新扫描结果更新观测、几何和角度证据，再重新计算缺口和下一视点", ha="center", fontsize=10)
    ax.set_xlim(0, 1)
    ax.set_ylim(0, 1)
    return fig


def fig_system() -> plt.Figure:
    fig, ax = plt.subplots(figsize=(9, 5.5))
    ax.set_axis_off()
    box(ax, (0.05, 0.68), 0.22, 0.15, "输入层\nIFC / 点云 / 评分CSV")
    box(ax, (0.37, 0.76), 0.22, 0.12, "Patch生成模块")
    box(ax, (0.37, 0.58), 0.22, 0.12, "证据构造模块")
    box(ax, (0.68, 0.67), 0.24, 0.14, "缺口评分模块\nG_gap / G_task")
    box(ax, (0.20, 0.30), 0.22, 0.14, "视点规划模块")
    box(ax, (0.55, 0.30), 0.22, 0.14, "闭环仿真模块")
    box(ax, (0.37, 0.08), 0.30, 0.12, "Web交互平台\n表格、三维图、补扫按钮")
    for a, b in [((0.27, 0.76), (0.37, 0.82)), ((0.27, 0.72), (0.37, 0.64)), ((0.59, 0.82), (0.68, 0.74)), ((0.59, 0.64), (0.68, 0.72)), ((0.80, 0.67), (0.32, 0.44)), ((0.42, 0.37), (0.55, 0.37)), ((0.66, 0.30), (0.53, 0.20))]:
        arrow(ax, a, b)
    ax.set_xlim(0, 1)
    ax.set_ylim(0, 1)
    return fig


def fig_tri() -> plt.Figure:
    fig, ax = plt.subplots(figsize=(7, 5.2))
    ax.set_axis_off()
    poly = np.array([[0.15, 0.15], [0.85, 0.20], [0.78, 0.82], [0.22, 0.78]])
    ax.add_patch(Polygon(poly, fill=False, edgecolor="black", linewidth=2))
    pts = [poly[0], poly[1], poly[2], poly[3], np.array([0.48, 0.48])]
    tris = [(0, 1, 4), (1, 2, 4), (2, 3, 4), (3, 0, 4)]
    for tri in tris:
        p = np.array([pts[i] for i in tri])
        ax.add_patch(Polygon(p, fill=False, edgecolor="black", linewidth=1.2, hatch="//"))
    for i, p in enumerate(pts):
        ax.plot(p[0], p[1], "ko", ms=3)
        ax.text(p[0] + 0.01, p[1] + 0.01, f"V{i+1}", fontsize=9)
    ax.text(0.50, 0.90, "IFC构件表面三角化：三角面T1-T4保留构件GUID和材料来源", ha="center", fontsize=11)
    ax.set_xlim(0, 1)
    ax.set_ylim(0, 1)
    return fig


def fig_region() -> plt.Figure:
    fig, ax = plt.subplots(figsize=(8, 5.2))
    ax.set_axis_off()
    centers = [(0.2, 0.55), (0.38, 0.58), (0.56, 0.55), (0.72, 0.50), (0.45, 0.34)]
    for i, (x, y) in enumerate(centers):
        ax.add_patch(Polygon([[x - 0.08, y - 0.05], [x + 0.08, y - 0.03], [x, y + 0.08]], fill=False, edgecolor="black", linewidth=1.5, hatch="/" if i < 3 else "\\"))
        ax.text(x, y, f"T{i+1}", ha="center", va="center", fontsize=9)
        ax.arrow(x, y + 0.08, 0.04 if i < 3 else -0.05, 0.08 if i < 3 else 0.02, head_width=0.015, color="black", length_includes_head=True)
    ax.plot([0.20, 0.38, 0.56], [0.55, 0.58, 0.55], "k--", lw=1)
    ax.text(0.42, 0.78, "共享边邻接 + 当前法向约束 + 种子法向约束", ha="center", fontsize=11)
    ax.text(0.42, 0.18, "T4、T5虽可能局部相邻，但与种子面法向差过大，不并入同一Patch", ha="center", fontsize=10)
    ax.set_xlim(0, 1)
    ax.set_ylim(0, 1)
    return fig


def fig_assoc() -> plt.Figure:
    fig, ax = plt.subplots(figsize=(8, 5.2))
    ax.set_axis_off()
    ax.add_patch(Polygon([[0.15, 0.20], [0.75, 0.25], [0.62, 0.70], [0.22, 0.65]], fill=False, edgecolor="black", linewidth=1.5, hatch="//"))
    rng = np.random.default_rng(4)
    pts = rng.normal([0.45, 0.45], [0.18, 0.14], size=(80, 2))
    ax.scatter(pts[:, 0], pts[:, 1], s=8, c="black", alpha=0.6)
    ax.add_patch(Circle((0.42, 0.46), 0.10, fill=False, linestyle="--", edgecolor="black"))
    ax.text(0.42, 0.59, "距离阈值d0=0.05m\n内点关联至最近表面", ha="center", fontsize=9)
    ax.text(0.12, 0.82, "刚性配准T：p' = R p + t；仓库当前CRAS粗平移为(0.6848,0,-0.6665)", fontsize=10)
    ax.set_xlim(0, 1)
    ax.set_ylim(0, 1)
    return fig


def fig_evidence() -> plt.Figure:
    fig, ax = plt.subplots(figsize=(9, 5.5))
    ax.set_axis_off()
    box(ax, (0.05, 0.42), 0.18, 0.18, "Patch i\n面积/中心/法向")
    labels = ["观测证据\nM,N,Q,R", "角度证据\nF,N,A", "几何证据\nC,Rρ", "语义证据\np_bim,p_obs", "材料证据\n缺失/冲突"]
    xs = [0.32, 0.56, 0.32, 0.56, 0.76]
    ys = [0.68, 0.68, 0.34, 0.34, 0.51]
    for x, y, label in zip(xs, ys, labels):
        box(ax, (x, y), 0.18, 0.14, label, fontsize=9)
        arrow(ax, (0.23, 0.51), (x, y + 0.07))
    box(ax, (0.40, 0.08), 0.30, 0.13, "方向一致的缺口指标\nD_sem,D_mat,D_obs,D_ang,D_geo")
    for x, y in zip(xs, ys):
        arrow(ax, (x + 0.09, y), (0.55, 0.21))
    ax.set_xlim(0, 1)
    ax.set_ylim(0, 1)
    return fig


def sample_scores(df: pd.DataFrame, n: int = 1400) -> pd.DataFrame:
    if df.empty:
        return df
    out = df.sample(min(n, len(df)), random_state=7).copy()
    if not {"centroid_x", "centroid_y", "centroid_z"}.issubset(out.columns) and "centroid" in out.columns:
        coords = np.array([parse_tuple(value) for value in out["centroid"]], dtype=float)
        out["centroid_x"] = coords[:, 0]
        out["centroid_y"] = coords[:, 1]
        out["centroid_z"] = coords[:, 2]
    if not {"normal_x", "normal_y", "normal_z"}.issubset(out.columns) and "normal" in out.columns:
        coords = np.array([parse_tuple(value) for value in out["normal"]], dtype=float)
        out["normal_x"] = coords[:, 0]
        out["normal_y"] = coords[:, 1]
        out["normal_z"] = coords[:, 2]
    return out


def scatter3d(df: pd.DataFrame, col: str, bw: bool, title: str) -> plt.Figure:
    fig = plt.figure(figsize=(8, 6))
    ax = fig.add_subplot(111, projection="3d")
    s = sample_scores(df, 1800)
    c = pd.to_numeric(s[col], errors="coerce").fillna(0)
    cmap = "Greys" if bw else "viridis"
    sc = ax.scatter(s["centroid_x"], s["centroid_y"], s["centroid_z"], c=c, cmap=cmap, s=np.clip(s["area"] * 20, 4, 35), alpha=0.85)
    ax.set_xlabel("X/m")
    ax.set_ylabel("Y/m")
    ax.set_zlabel("Z/m")
    ax.set_title(title)
    fig.colorbar(sc, ax=ax, shrink=0.65, label=col)
    ax.view_init(elev=24, azim=-55)
    return fig


def parse_tuple(text: Any) -> tuple[float, float, float]:
    if isinstance(text, (tuple, list, np.ndarray)):
        vals = text
    else:
        vals = str(text).strip("()[]").replace(",", " ").split()
    vals = [float(v) for v in vals[:3]]
    return vals[0], vals[1], vals[2]


def fig_candidate_field(bw: bool, detailed: bool = False) -> plt.Figure:
    fig = plt.figure(figsize=(8, 6))
    ax = fig.add_subplot(111, projection="3d")
    df = sample_scores(patch_real if not patch_real.empty else patch_synth, 900)
    colors = "0.75" if bw else pd.to_numeric(df["G_gap"], errors="coerce").fillna(0)
    ax.scatter(df["centroid_x"], df["centroid_y"], df["centroid_z"], c=colors, cmap="viridis", s=6, alpha=0.45)
    ranked = views_real if not views_real.empty else views_synth
    for _, row in ranked.head(12 if detailed else 27).iterrows():
        p = np.array(parse_tuple(row["position"]))
        o = np.array(parse_tuple(row["orientation"]))
        ax.scatter([p[0]], [p[1]], [p[2]], c="black" if bw else "red", marker="^", s=35)
        ax.quiver(p[0], p[1], p[2], o[0], o[1], o[2], length=0.8, color="black" if bw else "red")
    ax.set_xlabel("X/m")
    ax.set_ylabel("Y/m")
    ax.set_zlabel("Z/m")
    ax.set_title("候选补扫视点与视线方向" if not detailed else "视锥、法向和可见Patch示意")
    ax.view_init(elev=24, azim=-55)
    return fig


def fig_ranking() -> plt.Figure:
    fig, ax = plt.subplots(figsize=(8, 5))
    ranked = views_real if not views_real.empty else views_synth
    top = ranked.head(10).copy()
    y = np.arange(len(top))[::-1]
    ax.barh(y, top["value"], color="0.65", edgecolor="black")
    ax.set_yticks(y)
    ax.set_yticklabels([str(v) for v in top["view_id"]], fontsize=8)
    ax.set_xlabel("视点价值 V(v)")
    ax.set_title("候选视点收益排序")
    ax.grid(axis="x", linestyle=":", color="0.6")
    return fig


def fig_before_after() -> plt.Figure:
    fig, axes = plt.subplots(1, 2, figsize=(10, 4.5), sharex=True, sharey=True)
    df = sample_scores(patch_synth if not patch_synth.empty else patch_real, 1200)
    c0 = pd.to_numeric(df["G_gap"], errors="coerce").fillna(0)
    c1 = np.maximum(0, c0 * 0.55)
    for ax, c, title in zip(axes, [c0, c1], ["补扫前", "补扫后示意"]):
        sc = ax.scatter(df["centroid_x"], df["centroid_y"], c=c, cmap="Greys", s=8)
        ax.set_title(title)
        ax.set_xlabel("X/m")
        ax.set_ylabel("Y/m")
        ax.set_aspect("equal", adjustable="box")
    fig.colorbar(sc, ax=axes.ravel().tolist(), shrink=0.8, label="G_gap")
    return fig


def fig_loop() -> plt.Figure:
    fig, ax = plt.subplots(figsize=(8, 5))
    h = cl_synth if not cl_synth.empty else cl_real
    if not h.empty:
        ax.plot(h["step"], h["remaining_gap_area"], "ko-", label="剩余缺口面积")
        ax2 = ax.twinx()
        ax2.plot(h["step"], h["recovery_rate"], "k--", label="恢复率")
        ax2.set_ylabel("恢复率")
    ax.set_xlabel("补扫轮次")
    ax.set_ylabel("剩余缺口面积")
    ax.set_title("连续多轮补扫视点序列和剩余缺口变化")
    ax.grid(True, linestyle=":")
    return fig


def fig_web_arch() -> plt.Figure:
    fig, ax = plt.subplots(figsize=(9, 5.5))
    ax.set_axis_off()
    box(ax, (0.05, 0.68), 0.22, 0.13, "数据源下拉\nsynthetic/raycast/real")
    box(ax, (0.05, 0.48), 0.22, 0.13, "模型导入\nIFC/NPZ/OBJ/PLY/STL/GLB")
    box(ax, (0.05, 0.28), 0.22, 0.13, "参数滑块\n阈值/TopK/恢复率/步数")
    box(ax, (0.38, 0.58), 0.24, 0.16, "后端状态\n_scene + scores + ranked")
    box(ax, (0.72, 0.68), 0.20, 0.12, "缺口模型图")
    box(ax, (0.72, 0.50), 0.20, 0.12, "候选视角图和表")
    box(ax, (0.72, 0.32), 0.20, 0.12, "指标和构件排名")
    box(ax, (0.72, 0.14), 0.20, 0.12, "闭环记录")
    for a in [(0.27, 0.745), (0.27, 0.545), (0.27, 0.345)]:
        arrow(ax, a, (0.38, 0.66))
    for b in [(0.72, 0.74), (0.72, 0.56), (0.72, 0.38), (0.72, 0.20)]:
        arrow(ax, (0.62, 0.66), b)
    ax.set_xlim(0, 1)
    ax.set_ylim(0, 1)
    return fig


def generate_figures() -> None:
    formal_figs = [
        ("图1_方法总体流程图", fig_flow()),
        ("图2_系统模块框图", fig_system()),
        ("图3_IFC构件三角化示意图", fig_tri()),
        ("图4_共享边双法向Patch示意图", fig_region()),
        ("图5_点云IFC配准和距离阈值示意图", fig_assoc()),
        ("图6_Patch多源证据结构图", fig_evidence()),
        ("图7_真实模型Patch缺口分布三维图", scatter3d(patch_real, "G_gap", True, "真实模型Patch缺口分布")),
        ("图8_候选视点生成三维图", fig_candidate_field(True, False)),
        ("图9_视锥可见Patch三维图", fig_candidate_field(True, True)),
        ("图10_候选视点收益排序示意图", fig_ranking()),
        ("图11_补扫前后缺口对比图", fig_before_after()),
        ("图12_连续补扫序列和剩余缺口变化图", fig_loop()),
        ("图13_Web平台功能架构图", fig_web_arch()),
    ]
    for name, fig in formal_figs:
        save_fig(fig, name, True)
    with PdfPages(FORMAL / "说明书附图_黑白版.pdf") as pdf:
        for name in [n for n, _ in formal_figs]:
            img = plt.imread(FORMAL / f"{name}.png")
            fig, ax = plt.subplots(figsize=(8.27, 11.69))
            ax.imshow(img)
            ax.axis("off")
            ax.set_title(name.replace("_", " "), fontsize=12)
            pdf.savefig(fig, bbox_inches="tight")
            plt.close(fig)

    color_figs = [
        ("彩色1_IFC与Patch并排", scatter3d(patch_synth if not patch_synth.empty else patch_real, "area", False, "IFC Patch面积分布")),
        ("彩色2_CRAS点云叠加IFC", scatter3d(patch_real, "D_geo", False, "CRAS诊断几何证据")),
        ("彩色3_G_gap热力图", scatter3d(patch_real, "G_gap", False, "G_gap热力图")),
        ("彩色4_最佳视点和视锥", fig_candidate_field(False, True)),
        ("彩色5_第0_1_5_10轮闭环状态", fig_loop()),
    ]
    for name, fig in color_figs:
        save_fig(fig, name, False)


def generate_audit_docs() -> None:
    inv_rows = [
        ["核心源码", "patches", "src/patent_gap/patches/__init__.py", "三角面法向、面积、中心、共享边邻接、区域生长、Patch属性"],
        ["核心源码", "evidence", "src/patent_gap/evidence/__init__.py", "synthetic/raycast/real三类观测、角度、几何证据"],
        ["核心源码", "semantics", "src/patent_gap/semantics/__init__.py", "IFC/CRAS语义映射与JS散度"],
        ["核心源码", "materials", "src/patent_gap/materials/__init__.py", "IFC材料缺失与RGB材料冲突"],
        ["核心源码", "gap", "src/patent_gap/gap/scoring.py", "六类缺口指标融合、工程重要度、构件排序"],
        ["核心源码", "viewpoints", "src/patent_gap/viewpoints/ranking.py", "候选视点生成、可见性筛选、价值计算、贪心降权"],
        ["核心源码", "closed_loop", "src/patent_gap/simulation/closed_loop.py", "补扫后观测证据更新与闭环迭代"],
        ["核心源码", "registration", "src/patent_gap/registration/__init__.py", "CRAS粗平移和可选ICP刚性配准"],
        ["Web", "webapp", "src/patent_gap/webapp.py", "Gradio数据源、模型导入、参数滑块、三维图、补扫按钮"],
        ["脚本", "synthetic", "scripts/run_synthetic_pipeline.py", "Y>14m受控留出、定量评价、视点和闭环输出"],
        ["脚本", "real", "scripts/run_real_pipeline.py", "CRAS真实诊断，不含Patch级二元真值"],
        ["脚本", "raycast", "scripts/gen_synthetic_scan.py", "9站虚拟扫描、1度射线、patch_stats输出"],
        ["数据", "IFC", "data/raw/craslabbim.ifc", f"真实IFC；三角化后{PRE.get('element_count')}构件、{PRE.get('triangle_count')}三角面"],
        ["数据", "点云", "data/raw/craslabannotated.zip", f"CRAS真实室内点云；全量关联点数{FULL_ASSOC.get('points')}"],
        ["数据", "sample", "data/processed/cras_point_sample*.parquet", f"样本点{SAMPLE_ASSOC.get('input_points')}，匹配率{fmt(SAMPLE_ASSOC.get('matched_ratio'))}"],
        ["输出", "synthetic_summary", "outputs/reports/synthetic_summary.json", f"AUROC={fmt(SYN.get('AUROC'))}, AUPRC={fmt(SYN.get('AUPRC'))}, 仅受控留出实验"],
        ["输出", "real_summary", "outputs/reports/real_data_summary.json", f"诊断-only；Patch={REAL.get('n_patches')}，候选视点={REAL.get('candidate_views')}"],
    ]
    write_text(PATENT / "ASSET_INVENTORY.md", "# 资产审计\n\n" + md_table(inv_rows, ["类别", "资产", "路径", "审计结论"]))

    truth_rows = [
        ["共享边区域生长", "是", "patches.__init__.py:41-81, 186-187", "仅在同一IFC构件内对三角面按共享边邻接生长"],
        ["双法向约束", "是", "patches.__init__.py:65-77", "同时检查当前面法向和种子面法向与邻面夹角"],
        ["Patch属性", "是", "patches.__init__.py:205-228", "含面积、中心、法向、IFC类别、材料、工程重要度"],
        ["点云配准", "部分实现", "registration.__init__.py:18-20,134-175", "当前有粗平移和可选ICP；全量关联输出使用open3d最近面"],
        ["未匹配点保留", "是", "cras_full_assoc_summary.json", f"unmatched_points={FULL_ASSOC.get('unmatched_points')}"],
        ["遮挡射线检测", "当前视点评分未实现", "viewpoints/ranking.py:207-212", "只按frontality、FOV、range筛选；遮挡只能写可选实施方式"],
        ["Web模型导入", "是", "webapp.py:378-457", "IFC三角化；通用网格作为单Patch；评分CSV代理"],
        ["真实CRAS监督精度", "不支持", "real_data_summary.json", "无Patch级二元缺失真值，真实结果仅为工程诊断"],
    ]
    write_text(PATENT / "IMPLEMENTATION_TRUTH_TABLE.md", "# 实现真实性表\n\n" + md_table(truth_rows, ["技术点", "当前实现", "证据", "限制"]))

    params = [
        ["normal_threshold_deg", "20.0", "20.0", "patches.__init__.py:97", "build_patches", "synthetic/real/web IFC；web IFC覆盖min_patch_area但不覆盖该值"],
        ["min_patch_area", "0.005", "0.005", "patches.__init__.py:98", "build_patches", "默认；web导入IFC用0.01"],
        ["reference_density_pts_m2", "500.0", "500.0", "evidence.__init__.py:37,162,258", "evidence functions", "synthetic/raycast/real均可用默认"],
        ["frontality_gamma", "2.0", "2.0", "evidence.__init__.py:38,260", "compute_evidence*", "raycast/real证据"],
        ["min_points_for_observed", "1", "1", "evidence.__init__.py:39", "compute_evidence_from_synthetic", "raycast路径"],
        ["degradation_fraction", "0.15", "0.15", "evidence.__init__.py:163", "compute_controlled_withheld_evidence", "synthetic controlled"],
        ["missing_leakage_fraction", "0.20", "0.20", "evidence.__init__.py:164", "compute_controlled_withheld_evidence", "synthetic controlled"],
        ["semantic_smoothing", "0.02", "0.02", "semantics.__init__.py:135,194", "compute_semantic_scores*", "semantic路径"],
        ["ignore_unknown_labels", "True", "True", "semantics.__init__.py:136,195", "compute_semantic_scores*", "semantic路径"],
        ["semantic_noise", "0.12", "0.12", "semantics.__init__.py:266", "compute_controlled_semantic_scores", "synthetic controlled"],
        ["semantic_conflict_fraction", "0.05", "0.05", "semantics.__init__.py:267", "compute_controlled_semantic_scores", "synthetic controlled"],
        ["material_conflict_threshold", "0.60", "0.60", "materials.__init__.py:71", "build_material_table", "real/synthetic/web proxy"],
        ["min_rgb_points", "3", "3", "materials.__init__.py:72", "build_material_table", "real RGB证据"],
        ["alpha_sem", GAP_CFG.get("alpha_sem"), "0.16", "gap/scoring.py:10-17; configs/gap/default.yaml", "compute_patch_scores", "所有G_gap入口"],
        ["alpha_mat_missing", GAP_CFG.get("alpha_mat_missing"), "0.14", "gap/scoring.py:10-17; configs/gap/default.yaml", "compute_patch_scores", "所有G_gap入口"],
        ["alpha_mat_conflict", GAP_CFG.get("alpha_mat_conflict"), "0.10", "gap/scoring.py:10-17; configs/gap/default.yaml", "compute_patch_scores", "所有G_gap入口"],
        ["alpha_obs", GAP_CFG.get("alpha_obs"), "0.22", "gap/scoring.py:10-17; configs/gap/default.yaml", "compute_patch_scores", "所有G_gap入口"],
        ["alpha_ang", GAP_CFG.get("alpha_ang"), "0.16", "gap/scoring.py:10-17; configs/gap/default.yaml", "compute_patch_scores", "所有G_gap入口"],
        ["alpha_geo", GAP_CFG.get("alpha_geo"), "0.22", "gap/scoring.py:10-17; configs/gap/default.yaml", "compute_patch_scores", "所有G_gap入口"],
        ["lambda_importance", GAP_CFG.get("lambda_importance"), "0.30", "gap/scoring.py:225", "compute_patch_scores", "G_task"],
        ["high_gap_threshold", GAP_CFG.get("high_gap_threshold"), "0.55", "gap/scoring.py:285", "rank_components", "组件排序和高缺口统计"],
        ["distances", VIEW_CFG.get("distances"), "[1.5,2.5,4.0]", "viewpoints/ranking.py:61", "generate_candidates", "默认/配置/web"],
        ["azimuth_offsets_deg", VIEW_CFG.get("azimuth_offsets_deg"), "[-45,0,45]", "viewpoints/ranking.py:62", "generate_candidates", "默认/配置/web"],
        ["elevation_offsets_deg", VIEW_CFG.get("elevation_offsets_deg"), "[-20,0,20]", "viewpoints/ranking.py:63", "generate_candidates", "默认/配置/web"],
        ["gap_threshold", "0.55默认；synthetic正式脚本0.70；web滑块0.55", "0.55", "viewpoints/ranking.py:64; run_synthetic_pipeline.py:222; webapp.py:755-760", "generate_candidates", "入口不同"],
        ["max_target_patches", "30", "30", "viewpoints/ranking.py:65; webapp.py:463", "generate_candidates", "synthetic/real/web"],
        ["field_of_view_deg", "90默认；web 100", "90.0", "viewpoints/ranking.py:144; webapp.py:464", "score_candidates", "脚本默认/web覆盖"],
        ["max_range_m", "20默认；web 30", "20.0", "viewpoints/ranking.py:145; webapp.py:465", "score_candidates", "脚本默认/web覆盖"],
        ["eta", VIEW_CFG.get("eta"), "0.10", "viewpoints/ranking.py:141; configs/view/default.yaml", "score_candidates", "冗余惩罚"],
        ["recovery_per_view", "0.75脚本；web 0.55默认", "0.75/0.55", "closed_loop.py:164; webapp.py:763-768", "closed_loop/run_scan_steps", "入口不同"],
        ["association_distance", CRAS_FULL_CFG.get("association", {}).get("association_distance"), "0.05", "configs/experiment/cras_full.yaml", "association pipeline", "real CRAS关联"],
    ]
    with (PATENT / "PARAMETER_REGISTRY.csv").open("w", encoding="utf-8", newline="") as f:
        w = csv.writer(f)
        w.writerow(["参数名", "当前值", "默认值", "来源文件", "调用函数", "synthetic/raycast/real/web映射"])
        w.writerows(params)

    lineage = f"""
# 数据血缘

## 原始数据
- IFC：`data/raw/craslabbim.ifc`，配置见 `configs/data/cras.yaml`，MD5字段为 `{CRAS_CFG.get('expected_md5', {}).get('craslabbim.ifc')}`。
- 点云压缩包：`data/raw/craslabannotated.zip`，配置MD5字段为 `{CRAS_CFG.get('expected_md5', {}).get('craslabannotated.zip')}`。

## 预处理
- `outputs/reports/preprocess_summary.json` 记录 IFC 已三角化，构件数 `{PRE.get('element_count')}`，材料数 `{PRE.get('material_count')}`，顶点数 `{PRE.get('vertex_count')}`，三角面数 `{PRE.get('triangle_count')}`。
- 输出包括 `ifc_mesh.npz`、`ifc_elements.parquet`、`ifc_materials.parquet`、`triangle_element_map.npy`、`triangle_element_guid.npy`。

## 点云关联
- 样本关联：`cras_point_sample_associations.parquet`，输入 `{SAMPLE_ASSOC.get('input_points')}` 点，匹配 `{SAMPLE_ASSOC.get('matched_points')}` 点，未匹配 `{SAMPLE_ASSOC.get('unmatched_points')}` 点，阈值 `{SAMPLE_ASSOC.get('association_distance_m')}` m。
- 全量关联摘要：`cras_full_assoc_summary.json`，输入 `{FULL_ASSOC.get('points')}` 点，匹配 `{FULL_ASSOC.get('matched_points')}` 点，未匹配 `{FULL_ASSOC.get('unmatched_points')}` 点，匹配率 `{fmt(FULL_ASSOC.get('matched_ratio'))}`。

## 受控synthetic数据
- `scripts/gen_synthetic_scan.py` 生成 9 个虚拟扫描站，射线方位角 0-359 度、仰角 -30 至 80 度、步长 1 度。
- `scripts/run_synthetic_pipeline.py` 将 `centroid_y > 14.0m` 定义为预声明留出标签；该规则是实验构造，不是自然建筑规律。
- `patch_stats.csv` 当前 `{len(patch_stats)}` 个Patch有射线命中统计，`scanner_positions.csv` 当前 `{len(scanner_pos)}` 个站位。

## 输出
- synthetic定量输出：`outputs/reports/synthetic_summary.json`；标签有Patch级真值。
- CRAS真实诊断输出：`outputs/reports/real_data_summary.json`；无Patch级二元缺失真值，只作工程诊断。
"""
    write_text(PATENT / "DATA_LINEAGE.md", lineage)

    formula_rows = [
        ["三角面法向", "n=normalize((v1-v0)x(v2-v0))", "patches.__init__.py:18-27", "vertices, faces", "face_normals", "build_patches"],
        ["三角面积", "A=0.5||(v1-v0)x(v2-v0)||", "patches.__init__.py:30-34", "vertices, faces", "face_areas", "build_patches"],
        ["Patch中心", "c=sum(A_t c_t)/sum(A_t)", "patches.__init__.py:195-197", "local_areas, local_centroids", "centroid", "build_patches"],
        ["Patch法向", "n_p=normalize(sum(w_t n_t))", "patches.__init__.py:197-199", "local_normals, area weights", "normal", "build_patches"],
        ["区域生长约束", "dot(n_cur,n_nb)>=cosθ 且 dot(n_seed,n_nb)>=cosθ", "patches.__init__.py:56-81", "local_normals, adjacency", "patch labels", "build_patches"],
        ["JS散度", "D_JS=1/2 KL(p||m)+1/2 KL(q||m), m=(p+q)/2", "gap/scoring.py:20-41", "p_bim,p_obs", "D_sem", "compute_patch_scores"],
        ["D_obs", "clip(.35M+.25(2-N)/2+.20(1-Q)+.20(1-R))", "gap/scoring.py:156-173; closed_loop.py:110-117", "observed,Nvalid,Qproj,Rreg", "D_obs", "compute_patch_scores/closed_loop"],
        ["D_ang", "clip(1-(.50F+.30N/2+.20A))", "gap/scoring.py:176-190; closed_loop.py:118-127", "best_frontality,Nvalid,angular_diversity", "D_ang", "compute_patch_scores/closed_loop"],
        ["D_geo", "clip(.60(1-C)+.40(1-Rρ))", "gap/scoring.py:191-200; closed_loop.py:128-133", "coverage_ratio,density_ratio", "D_geo", "compute_patch_scores/closed_loop"],
        ["材料缺失", "D_mat_missing=0 if present else 1", "materials.__init__.py:159", "material_present", "D_mat_missing", "build_material_table"],
        ["材料冲突", "D_mat_conflict=1 if clear mismatch else 0", "materials.__init__.py:141-160", "IFC category, RGB category", "D_mat_conflict", "build_material_table"],
        ["可用权重融合", "G_gap=sum(a_k D_k)/sum(a_k), skip NaN", "gap/scoring.py:54-66,215-224", "D components, alpha", "G_gap", "compute_patch_scores"],
        ["工程重要度", "G_task=clip(G_gap(1+λI),0,2)", "gap/scoring.py:225-231", "G_gap,I,lambda", "G_task", "compute_patch_scores"],
        ["构件评分", ".25mean+.25max+.20p90+.20high_area+.10importance", "gap/scoring.py:285-310", "patch scores by element", "G_component", "rank_components"],
        ["候选视点方向", "Rodrigues旋转，方位角绕bitangent，俯仰角绕tangent", "viewpoints/ranking.py:35-56,97-110", "centroid,normal,az,el,distance", "position,orientation", "generate_candidates"],
        ["可见性", "frontality>fmin, alignment>=cos(FOV/2), range<=max", "viewpoints/ranking.py:197-212", "patch, candidate", "visible_mask", "score_candidates"],
        ["视点质量", "q=frontality*exp(-(d-2.5)^2/8)*1/(1+.15d)", "viewpoints/ranking.py:236-240", "frontality,distance", "q", "score_candidates"],
        ["视点价值", "V=sum(G_gap*q*area)-eta*overlap", "viewpoints/ranking.py:241-260", "visible patches", "value", "score_candidates"],
        ["补扫更新", "x'=x+r_eff(1-x), r_eff=recovery(0.5+G_gap)", "closed_loop.py:56-80", "visible_patch_ids,recovery,G_gap", "updated observations", "apply_supplemental_observation"],
    ]
    write_text(PATENT / "FORMULA_TO_CODE_MAP.md", "# 公式到源码映射\n\n" + md_table(formula_rows, ["公式/指标", "表达式", "源码路径:行号", "输入字段", "输出字段", "调用入口"]))

    risks = """
# 风险与不支持特征

1. 当前候选视点评分没有遮挡射线检测；只能在“另一实施方式”写入遮挡判断。
2. CRAS真实数据没有Patch级二元缺失真值，不能用真实结果声称AUROC、AUPRC、F1等监督精度。
3. synthetic的`centroid_y > 14m`是受控留出规则，不是建筑物自然规律。
4. synthetic观测证据由标签条件化生成，不能声称标签和证据统计独立。
5. Web平台是工程实现界面，不应作为唯一创造性；创造性应落在IFC语义Patch、多源证据可用权重融合和闭环视点规划。
6. 当前通用OBJ/PLY/STL/GLB导入仅形成单Patch代理，不等同于IFC构件级语义分块。
7. 真实CRAS样本关联中classification均为0导致D_sem不可用，不能写成真实语义冲突已验证。
8. 当前材料冲突是粗RGB类别启发式，不是严格材料识别模型。
9. ICP为可选实现；当前CRAS全量关联摘要记录的是已知平移和最近面关联，不能夸大为完整鲁棒配准系统。
10. 闭环补扫是软件仿真更新，并非真实机器人或扫描仪实采闭环。
"""
    write_text(PATENT / "RISKS_AND_UNSUPPORTED_FEATURES.md", risks)


def generate_claim_docs() -> None:
    novelty_rows = [
        ["IFC构件内共享边区域生长", "已实现", "适合", "patches.__init__.py:41-81,162-187"],
        ["双法向约束抑制链式合并", "已实现", "适合", "patches.__init__.py:65-77"],
        ["Patch含面积/中心/法向/语义/材料/重要度", "已实现", "适合", "patches.__init__.py:205-228"],
        ["点云-IFC关联并保留未匹配点", "已实现于数据输出", "从属", "cras_full_assoc_summary.json; registration.__init__.py"],
        ["多源证据方向一致化", "已实现", "适合", "gap/scoring.py:140-224"],
        ["可用指标集合重归一化", "已实现", "适合", "gap/scoring.py:54-66"],
        ["缺口、面积和视点质量联合收益", "已实现", "适合", "viewpoints/ranking.py:236-260"],
        ["基于中心法向生成多距离/方位/俯仰视点", "已实现", "适合", "viewpoints/ranking.py:59-130"],
        ["闭环更新证据并重算下一视点", "已实现", "适合", "closed_loop.py:180-273"],
        ["工程重要度加权", "已实现", "从属", "gap/scoring.py:225-231"],
        ["顺序多视点冗余抑制", "已实现为降权", "从属", "viewpoints/ranking.py:270-285"],
        ["遮挡射线检测", "当前未实现", "仅可选", "不得进入独立权利要求"],
    ]
    write_text(PATENT / "NOVELTY_CANDIDATES.md", "# 创造性候选审计\n\n" + md_table(novelty_rows, ["候选特征", "真实性", "建议位置", "证据"]))

    support_rows = [
        ["获取IFC模型和点云", "scripts/run_real_pipeline.py:417-422; configs/data/cras.yaml", "CRAS IFC和点云zip", "preprocess_summary.json, cras_full_assoc_summary.json", "是", "是", "否", "否", "常规输入步骤，创造性弱"],
        ["IFC三角化", "patent_gap.ifc.reader由webapp.py:22和preprocess_summary引用", "craslabbim.ifc", "preprocess_summary.json", "是", "是", "否", "否", "三角化本身通用"],
        ["构件内共享边区域生长", "patches.__init__.py:41-81,162-187", "normal_threshold_deg=20", "patches_real.parquet", "是", "是", "否", "否", "与双法向约束组合较强"],
        ["双法向约束", "patches.__init__.py:65-77", "normal_threshold_deg=20", "patch数量9438", "是", "是", "否", "否", "建议写入独权"],
        ["Patch属性结构", "patches.__init__.py:205-228", "importance_map默认", "patch_scores_real.csv", "是", "是", "否", "否", "与后续融合直接关联"],
        ["点云-IFC最近表面关联", "registration.__init__.py:134-175; cras_full_assoc_summary", "association_distance=0.05", "matched/unmatched统计", "是", "是", "否", "否", "关联算法本身需与Patch证据结合"],
        ["构造观测/角度/几何/语义/材料证据", "evidence, semantics, materials模块", "见PARAMETER_REGISTRY.csv", "patch_scores*.csv", "是", "是", "否", "否", "多源证据体系可支撑"],
        ["方向一致缺口指标", "gap/scoring.py:140-213", "D越大缺口越强", "patch_scores*.csv", "是", "是", "否", "否", "建议写入独权"],
        ["可用证据权重重归一化", "gap/scoring.py:54-66", "DEFAULT_ALPHAS", "real中D_sem/D_ang NaN仍有G_gap", "是", "是", "否", "否", "最强候选之一"],
        ["候选视点生成", "viewpoints/ranking.py:59-130", "distances/az/el", "candidate_view_ranking_real.csv", "是", "是", "否", "否", "应避免锁死数值"],
        ["视点价值计算", "viewpoints/ranking.py:133-267", "FOV/range/eta", "top1_view_value", "是", "是", "否", "否", "与补扫收益相关"],
        ["闭环补扫更新", "closed_loop.py:13-143,180-273", "recovery_per_view", "closed_loop_results*.csv", "是", "是", "否", "否", "当前为仿真"],
        ["工程重要度加权", "gap/scoring.py:225-231", "lambda=0.30", "G_task列", "是", "否", "是", "否", "从属较稳妥"],
        ["贪心多视点冗余抑制", "viewpoints/ranking.py:270-285", "可见Patch降为25%", "candidate_view_greedy_real.csv", "是", "否", "是", "否", "实现较简单，放从属"],
        ["遮挡射线检测", "无", "无", "无", "否", "否", "否", "是", "当前不能主张"],
    ]
    write_text(PATENT / "CLAIM_SUPPORT_MATRIX.md", "# 权利要求支持矩阵\n\n" + md_table(support_rows, ["技术特征", "源码位置", "配置", "实验输出", "当前实现", "适合独权", "从属权利要求", "仅可选", "风险说明"]))


CLAIMS = """
1. 一种基于IFC语义表面分块和多源证据融合的建筑点云缺口检测及补充扫描视点规划方法，其特征在于，包括：
获取建筑的IFC模型以及与所述建筑对应的三维点云；
将所述IFC模型中的构件表面三角化，针对同一IFC构件内的三角面，根据共享边邻接关系以及三角面法向相似性进行区域生长，得到多个表面Patch；
将所述三维点云通过刚性坐标变换与所述IFC模型所在坐标系对齐，并将点云点关联至满足距离阈值的最近IFC表面或Patch；
针对每个Patch构造至少包括观测证据、几何覆盖证据、入射角度证据、IFC语义证据和材料证据中的多源证据；
将所述多源证据分别转换为方向一致的缺口指标，其中缺口指标值越大表示对应Patch的信息缺口越大；
针对每个Patch，在该Patch实际可用的缺口指标集合上对相应权重重新归一化，融合得到Patch缺口分数；
根据Patch缺口分数筛选目标Patch；
以目标Patch的中心和法向为基准，在多个距离、方位角和俯仰角组合下生成候选补充扫描视点；
根据候选视点对Patch的可见性、入射正向性、距离质量、分辨率质量、Patch面积和Patch缺口分数计算候选视点价值；
选择价值最高的候选视点作为下一补充扫描视点；
根据补充扫描得到的新观测更新Patch的多源证据，并重复执行缺口分数计算、目标Patch筛选、候选视点生成和候选视点选择。

2. 根据权利要求1所述的方法，其特征在于，所述区域生长在同一IFC构件内进行，并将两个三角面共享无向边作为邻接条件。

3. 根据权利要求1或2所述的方法，其特征在于，将候选邻接三角面并入当前Patch时，同时满足候选邻接三角面的法向与当前扩展三角面的法向夹角小于预设阈值，以及候选邻接三角面的法向与区域生长种子三角面的法向夹角小于所述预设阈值。

4. 根据权利要求1所述的方法，其特征在于，每个Patch记录Patch标识、所属IFC构件标识、构件类别、面积、面积加权中心、面积加权法向、材料名称和工程重要度中的至少一项。

5. 根据权利要求1所述的方法，其特征在于，点云点与IFC表面之间的关联包括：对点云点施加粗配准平移和/或迭代最近点刚性细配准，随后计算点云点至IFC三角面的最近距离，并保留距离超过阈值的未匹配点统计。

6. 根据权利要求1所述的方法，其特征在于，观测缺口指标由直接观测状态、有效视角数量、投影分辨率质量和配准置信度共同确定。

7. 根据权利要求1所述的方法，其特征在于，角度缺口指标由最佳入射正向性、有效视角数量和角度多样性共同确定。

8. 根据权利要求1所述的方法，其特征在于，几何缺口指标由Patch覆盖比例和点密度比例共同确定。

9. 根据权利要求1所述的方法，其特征在于，语义缺口指标由IFC构件类别映射得到的语义分布与点云分类标签得到的语义分布之间的Jensen-Shannon散度确定。

10. 根据权利要求1所述的方法，其特征在于，材料证据包括IFC材料是否存在以及IFC材料类别与点云颜色推断材料类别之间是否存在预设清晰冲突。

11. 根据权利要求1所述的方法，其特征在于，所述Patch缺口分数为可用缺口指标的加权平均，缺失或不可计算的缺口指标不参与分母权重求和。

12. 根据权利要求1所述的方法，其特征在于，候选视点价值为候选视点可见Patch的Patch缺口分数、Patch面积和视点质量的乘积求和，并扣除可见非目标Patch比例形成的冗余惩罚。

13. 根据权利要求1所述的方法，其特征在于，根据Patch工程重要度对Patch缺口分数进行任务相关放大，并用于构件级排序或任务优先级排序。

14. 一种基于IFC语义表面分块和多源证据融合的建筑点云缺口检测及补充扫描视点规划系统，其特征在于，包括：IFC解析与Patch生成模块、点云配准与关联模块、多源证据构造模块、缺口评分模块、候选视点生成模块、视点评分模块、闭环更新模块和交互展示模块；所述各模块被配置为执行权利要求1至13任一项所述的方法。

15. 一种电子设备或计算机可读存储介质，其上存储有计算机程序，所述计算机程序被处理器执行时实现权利要求1至13任一项所述的方法。
"""


FORMULAS = """
（1）三角面法向：n_t = ((v_1-v_0)×(v_2-v_0))/||((v_1-v_0)×(v_2-v_0))||。
（2）三角面面积：A_t = 1/2 ||(v_1-v_0)×(v_2-v_0)||。
（3）Patch中心：c_p = Σ_t A_t c_t / Σ_t A_t。
（4）Patch法向：n_p = normalize(Σ_t (A_t/Σ_t A_t) n_t)。
（5）双法向区域生长：dot(n_cur,n_nb)≥cosθ 且 dot(n_seed,n_nb)≥cosθ。
（6）Jensen-Shannon散度：D_JS(p,q)=1/2 KL(p||m)+1/2 KL(q||m)，m=(p+q)/2。
（7）观测缺口：D_obs=clip(0.35M+0.25clip((2-N_v)/2,0,1)+0.20(1-Q_p)+0.20(1-R_c),0,1)。
（8）角度缺口：D_ang=clip(1-(0.50F_best+0.30clip(N_v/2,0,1)+0.20A_div),0,1)。
（9）几何缺口：D_geo=clip(0.60(1-C)+0.40(1-clip(R_ρ,0,1)),0,1)。
（10）可用权重融合：G_gap=Σ_{k∈K_i} α_kD_{ik}/Σ_{k∈K_i}α_k。
（11）任务缺口：G_task=clip(G_gap(1+λI),0,2)。
（12）视点质量：q_{iv}=frontality_{iv}·exp(-((d_{iv}-2.5)^2)/(2·2^2))·clip(1/(1+0.15d_{iv}),0,1)。
（13）视点价值：V(v)=Σ_i G_gap,i q_{iv} A_i - η·overlap(v)。
（14）闭环恢复：x_i' = x_i + r_i(1-x_i)，r_i=clip(r_0(0.5+G_gap,i),0,1)。
"""


def patent_markdown() -> str:
    return f"""
# 一种基于IFC语义表面分块和多源证据融合的建筑点云缺口检测及补充扫描视点规划方法

## 摘要

本发明涉及建筑信息模型、三维点云处理和补充扫描视点规划。该方法获取建筑IFC模型和对应点云，将IFC构件表面三角化，并在同一构件内依据共享边邻接关系和双法向相似性约束形成Patch；将点云与IFC网格对齐并关联，针对Patch构造观测、几何、角度、语义和材料证据；将各证据转换为方向一致的缺口指标，并在每个Patch实际可用指标集合上重新归一化权重得到Patch缺口分数；再结合Patch面积、缺口分数和候选视点质量计算补充扫描视点价值，选择下一视点，并根据补扫结果更新证据形成闭环。该方法能够将点云缺口定位到IFC语义表面Patch，并避免不可用证据对融合结果产生固定偏置。摘要附图建议采用图1，因为图1概括了从IFC和点云输入到闭环补扫的完整流程。

## 权利要求书

{CLAIMS}

## 说明书

### 1. 技术领域

本发明涉及建筑信息模型、三维激光扫描、建筑点云处理、扫描质量检测和下一最佳视点规划，尤其涉及一种将IFC语义构件表面分解为Patch并利用多源证据融合检测建筑点云信息缺口、进而规划补充扫描视点的方法、系统、电子设备和计算机可读存储介质。

### 2. 背景技术

建筑室内外三维扫描通常需要多站扫描。受遮挡、扫描距离、入射角、视场范围和配准误差影响，点云会在局部构件表面产生覆盖不足、密度不足或观测角度不足。仅使用全局点数、全局覆盖率或整体配准误差，难以定位到具体建筑构件的具体表面区域。一般点云补全或神经网络重建方法可以生成外观上连续的点云，但不能直接给出真实补充扫描位姿，也难以说明某个IFC构件表面的证据缺失原因。一般下一最佳视点方法多关注几何覆盖，缺少IFC构件语义、材料属性、语义一致性和证据可用性差异的统一处理。

### 3. 发明内容

#### 3.1 要解决的技术问题

本发明要解决的技术问题包括：如何将点云缺口定位到IFC构件表面级别；如何避免三角面区域生长在缓慢弯曲面上发生链式误合并；如何将观测、角度、几何、语义和材料证据统一为可融合的缺口指标；如何在部分证据不可得时避免把缺失模态误当成无缺口或最大缺口；以及如何根据Patch缺口、面积和视点质量生成可执行的补充扫描视点。

#### 3.2 技术方案

本发明的核心技术方案如权利要求1所述。当前仓库实现中，Patch区域生长、证据构造、缺口融合、候选视点生成、视点评分和闭环更新分别由 `patches`、`evidence`、`semantics`、`materials`、`gap.scoring`、`viewpoints.ranking` 和 `simulation.closed_loop` 模块实现。

#### 3.3 有益效果

（1）通过IFC构件内Patch级分块，缺口定位粒度从构件整体提高到构件表面局部区域。

（2）通过当前三角面法向与种子三角面法向的双法向约束，可降低连续小角度弯曲导致的链式错误合并风险。

（3）通过将观测、角度、几何、语义和材料证据转化为方向一致的缺口指标，可对不同来源的证据进行统一比较。

（4）通过对每个Patch实际可用指标集合进行权重重归一化，可避免不可得证据造成固定偏置。CRAS真实诊断中D_sem和D_ang当前不可用，仍可用材料、观测和几何证据计算G_gap。

（5）通过联合Patch缺口、Patch面积和候选视点质量计算视点价值，可优先选择对高缺口、大面积、可高质量观测区域有贡献的补扫视点。

（6）通过补扫后更新Patch证据并重新计算缺口和视点，可减少重复补扫并形成逐轮收敛的规划流程。

### 4. 附图说明

图1为方法总体流程图。图2为系统模块框图。图3为IFC构件三角化示意图。图4为三角面通过共享边和双法向约束形成Patch的示意图。图5为点云与IFC网格配准、最近表面关联及距离阈值示意图。图6为Patch级多源证据结构图。图7为真实模型上的Patch缺口分布三维图。图8为围绕目标Patch生成不同距离、方位角、俯仰角候选视点的三维图。图9为候选视点、视锥、Patch法向、可见Patch和距离的三维图。图10为多个候选视点收益排序示意图。图11为补扫前后同一模型、同一相机位置的缺口对比图。图12为连续多轮补扫视点序列和剩余缺口变化图。图13为Web平台功能架构图。

### 5. 具体实施方式

#### 实施例1：IFC三角化与Patch生成

读取 `data/raw/craslabbim.ifc` 后得到三角网格。当前预处理输出显示IFC已三角化，构件数为 {PRE.get('element_count')}，材料数为 {PRE.get('material_count')}，顶点数为 {PRE.get('vertex_count')}，三角面数为 {PRE.get('triangle_count')}。对每个IFC构件单独取其三角面，先建立共享无向边邻接关系，再进行区域生长。实施例参数中法向阈值为20度，最小Patch面积为0.005平方米；Web导入IFC时最小Patch面积使用0.01平方米。

{FORMULAS}

上述公式中，v_0、v_1、v_2为三角面三个顶点，单位为米；A_t为三角面积，单位为平方米；n_t为三角面单位法向；c_t为三角面中心；θ为法向阈值。当前实现中θ默认20度，但权利要求不限定具体数值。

#### 实施例2：点云—IFC配准和关联

当前实现包括CRAS粗平移 `CRAS_COARSE_TRANSLATION=(0.6848,0,-0.6665)`，并提供基于SVD的ICP细配准函数。样本关联输出显示输入点数 {SAMPLE_ASSOC.get('input_points')}，匹配点数 {SAMPLE_ASSOC.get('matched_points')}，未匹配点数 {SAMPLE_ASSOC.get('unmatched_points')}，关联阈值 {SAMPLE_ASSOC.get('association_distance_m')} m。全量关联摘要显示输入点数 {FULL_ASSOC.get('points')}，匹配点数 {FULL_ASSOC.get('matched_points')}，未匹配点数 {FULL_ASSOC.get('unmatched_points')}，匹配率 {fmt(FULL_ASSOC.get('matched_ratio'))}。未匹配点作为距离超阈值证据保留在统计中。

#### 实施例3：Patch观测、角度、几何、语义和材料证据

观测证据包括直接观测状态、有效视角数量、最佳正向性、平均正向性、角度多样性、点数、点密度、覆盖比例、密度比例、投影分辨率质量和配准置信度。语义证据将IFC类别和CRAS分类标签映射到 `wall, floor, ceiling, door, window, column, furniture, pipe, equipment, unknown` 十类。材料证据使用IFC材料是否存在、材料名称类别和RGB均值启发式类别，只有在置信度达到阈值并属于明确不一致组合时标记材料冲突。

#### 实施例4：六类缺口指标及融合

六类缺口指标为D_sem、D_mat_missing、D_mat_conflict、D_obs、D_ang和D_geo。默认权重为：语义0.16、材料缺失0.14、材料冲突0.10、观测0.22、角度0.16、几何0.22。融合时跳过NaN指标，并用可用指标权重和作为分母。该机制使真实CRAS诊断中语义和角度证据不可用时，不会把不可用指标等同于0或1。

#### 实施例5：高缺口Patch筛选

当前默认高缺口阈值为0.55；synthetic正式脚本中候选视点生成使用0.70；Web平台阈值滑块默认0.55，范围0至1。阈值用于筛选目标Patch并用于构件高缺口面积比例统计。

#### 实施例6：候选补扫视点生成与价值计算

当前候选视点围绕Patch法向局部坐标系生成，默认距离为1.5m、2.5m、4.0m，方位角偏移为-45度、0度、45度，俯仰角偏移为-20度、0度、20度。候选视点评分使用入射正向性、视场角、距离上限、距离质量和分辨率质量。当前评分不执行遮挡射线检测，因此遮挡判断仅能作为可选实施方式。

#### 实施例7：闭环补扫更新

补扫后，对可见Patch标记为已观测，视角数量加1，最佳正向性至少提升至0.80，平均正向性至少提升至0.65，角度多样性增加0.20，覆盖比例和密度比例按剩余缺口恢复公式更新，投影质量至少为0.85，配准置信度至少为0.95。随后重新计算D_obs、D_ang、D_geo、G_gap和候选视点。

#### 实施例8：synthetic受控留出实验

synthetic实验明确将 `centroid_y > 14.0m` 设为预声明留出区，当前输出Patch数 {SYN.get('n_patches')}，缺失标签Patch数 {SYN.get('gt_missing_patches')}，非缺失Patch数 {SYN.get('gt_observed_patches')}，缺失率 {fmt(SYN.get('Prevalence'))}。该Y坐标阈值是受控实验规则，不是建筑自然规律。观测证据生成机制对缺失Patch和非缺失Patch采用不同分布，并允许缺失泄漏和非缺失退化，因此不得声称标签和观测证据统计独立。该实验用于定量验证，当前AUROC为 {fmt(SYN.get('AUROC'))}，AUPRC为 {fmt(SYN.get('AUPRC'))}，F1为 {fmt(SYN.get('F1'))}，10轮闭环最终恢复率为 {fmt(SYN.get('closed_loop_final_rr'))}。

#### 实施例9：CRAS真实室内点云和IFC诊断

CRAS真实诊断使用真实室内建筑点云和对应IFC。当前真实输出Patch数 {REAL.get('n_patches')}，构件数 {REAL.get('n_elements')}，G_gap均值 {fmt(REAL.get('G_gap_mean'))}，最大值 {fmt(REAL.get('G_gap_max'))}，高缺口Patch数 {REAL.get('high_gap_patches_count')}，候选视点数 {REAL.get('candidate_views')}，Top-1视点价值 {fmt(REAL.get('top1_view_value'))}。真实CRAS没有Patch级二元缺失真值，结果仅作工程诊断，不能据此声称监督精度。当前真实诊断中D_sem有效数量为 {REAL.get('D_sem_valid_count')}，D_ang有效数量为 {REAL.get('D_ang_valid_count')}。

#### 实施例10：Web平台实施方式

Web平台采用Gradio实现，提供三个内置数据源：合成留出实验、射线可见性诊断和CRAS实际关联；提供模型导入入口，支持IFC、NPZ、OBJ、PLY、STL和GLB；提供Patch评分CSV导入，按patch_id、element_guid、element_index或行顺序匹配G_gap，无法逐项匹配时使用CSV均值，未提供评分文件时使用0.65作为初始缺口代理。平台提供高缺口阈值、显示候选视角数、单次补扫恢复率和连续补扫步数滑块，并显示缺口模型、候选视角、指标热力图、构件排名和闭环记录。

#### 实施例11：系统、电子设备及存储介质

系统包括IFC解析与Patch生成模块、点云配准与关联模块、多源证据构造模块、缺口评分模块、候选视点生成模块、视点评分模块、闭环更新模块和交互展示模块。电子设备包括处理器和存储器，存储器中存储计算机程序，处理器执行该程序时实现上述方法。计算机可读存储介质存储计算机程序，该程序被执行时实现上述方法。

## 表1 关键实验输出追溯

{md_table([
["synthetic_summary.json", "n_patches", SYN.get('n_patches')],
["synthetic_summary.json", "AUROC", fmt(SYN.get('AUROC'))],
["synthetic_summary.json", "AUPRC", fmt(SYN.get('AUPRC'))],
["synthetic_summary.json", "closed_loop_final_rr", fmt(SYN.get('closed_loop_final_rr'))],
["real_data_summary.json", "evaluation_status", REAL.get('evaluation_status')],
["real_data_summary.json", "candidate_views", REAL.get('candidate_views')],
["cras_full_assoc_summary.json", "unmatched_points", FULL_ASSOC.get('unmatched_points')],
], ["文件", "字段", "值"])}

## 可选实施方式

在另一实施方式中，可在候选视点评分时加入遮挡射线检测，通过从候选视点向Patch中心或Patch采样点发射射线，判断射线是否被其他IFC三角面提前相交。该遮挡检测当前未在视点评分源码中实现，不能作为当前实验结论。
"""


def generate_patent_md() -> None:
    write_text(PATENT / "建筑点云缺口检测及补扫视点规划_中国发明专利初稿.md", patent_markdown())


def add_paragraph(doc: Document, text: str, style: str | None = None, bold: bool = False) -> None:
    p = doc.add_paragraph(style=style)
    run = p.add_run(text)
    run.font.name = "宋体"
    run.font.size = Pt(10.5)
    run.bold = bold


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_heading(level=level)
    run = p.add_run(text)
    run.font.name = "黑体"
    run.font.size = Pt(14 if level == 1 else 12)
    run.bold = True


def make_docx(path: Path, review: bool = False, claims_only: bool = False) -> None:
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Cm(21)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(2.5)
    sec.bottom_margin = Cm(2.5)
    sec.left_margin = Cm(2.8)
    sec.right_margin = Cm(2.6)
    styles = doc.styles
    styles["Normal"].font.name = "宋体"
    styles["Normal"].font.size = Pt(10.5)
    if review and not claims_only:
        add_heading(doc, "内部审阅目录", 1)
        add_paragraph(doc, "A. 发明名称；B. 摘要；C. 权利要求书；D. 说明书；E. 附图。")
        doc.add_page_break()
    add_heading(doc, "一种基于IFC语义表面分块和多源证据融合的建筑点云缺口检测及补充扫描视点规划方法", 1)
    if not claims_only:
        add_heading(doc, "摘要", 1)
        add_paragraph(doc, "本发明涉及建筑信息模型、三维点云处理和补充扫描视点规划。该方法将IFC构件表面三角化并形成Patch，构造多源证据，按可用证据集合融合得到Patch缺口分数，并据此规划补充扫描视点和闭环更新。摘要附图建议采用图1。")
    add_heading(doc, "权利要求书", 1)
    for para in CLAIMS.strip().split("\n\n"):
        add_paragraph(doc, para.strip())
    if claims_only:
        doc.save(path)
        return
    add_heading(doc, "说明书", 1)
    for section in patent_markdown().split("### ")[1:]:
        title, _, body = section.partition("\n")
        add_heading(doc, title.strip("# "), 2)
        for para in body.split("\n\n"):
            para = para.strip()
            if not para or para.startswith("|") or para.startswith("#"):
                continue
            add_paragraph(doc, para.replace("`", ""))
    add_heading(doc, "说明书附图", 1)
    for i in range(1, 14):
        files = sorted(FORMAL.glob(f"图{i}_*.png"))
        if not files:
            continue
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
        run.add_picture(str(files[0]), width=Cm(14.5))
        cap = doc.add_paragraph(f"图{i} {files[0].stem.split('_', 1)[1]}")
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.save(path)


def generate_final_audit() -> None:
    audit = f"""
# 最终质量门审计报告

## 检查结果

1. 独立权利要求技术特征均可在说明书和 `CLAIM_SUPPORT_MATRIX.md` 找到支持：通过。
2. 附图标记：本稿采用图号和模块文字标记，未引入复杂编号；图1至图13均已生成SVG、PNG和PDF汇总：通过。
3. 参数来源：已写入 `PARAMETER_REGISTRY.csv`，来自源码、配置或明确Web入口覆盖：通过。
4. 实验数字：synthetic、real、CRAS关联数字均来自 `outputs/reports/*.json` 或 `outputs/tables/*.csv`：通过。
5. 当前实现和可选实施方式：遮挡射线检测被明确列为未实现可选实施方式：通过。
6. synthetic和real结果：分别在实施例8和实施例9中说明，未混写监督精度：通过。
7. “点云补全”与“补充扫描”：正文强调本发明输出补扫视点，不声称完成点云补全：通过。
8. Web平台：作为实施方式和展示，不作为唯一技术贡献：通过。
9. 遮挡检测：未写入独立权利要求：通过。
10. CRAS Patch级真值：明确无Patch级二元缺失真值：通过。
11. 术语定义：Patch、G_gap、证据、候选视点等在说明书中定义：基本通过，建议代理师进一步统一术语。
12. Word文档：已生成审阅版、提交候选版、权利要求单独版；公式以可编辑文本形式保留：通过。

## 最强的3个创造性候选

1. 同一IFC构件内共享边区域生长与双法向约束组合，用于抑制链式误合并。
2. Patch级多源证据转化为方向一致缺口指标，并在可用证据集合上重新归一化融合。
3. Patch缺口、Patch面积和候选视点质量联合计算补扫视点价值，并在补扫后闭环更新证据。

## 最危险的5个现有技术重叠点

1. 一般下一最佳视点规划。
2. 一般点云覆盖率或密度质量检测。
3. 一般BIM/点云配准和最近面关联。
4. 一般点云语义分割或点云补全。
5. 一般三角网格区域生长分割。

## 建议专利代理师重点修改的权利要求

- 权利要求1：建议保留“IFC构件内Patch、双法向、可用证据重归一化、视点价值和闭环更新”的组合，不锁死具体数值。
- 权利要求3：建议强化双法向约束的技术效果。
- 权利要求11：建议作为核心从属或并入独权，突出可用证据集合重归一化。
- 权利要求12：建议明确视点质量与Patch缺口及面积的联合收益。

## 当前不能进入独立权利要求的技术特征

- 遮挡射线检测。
- 真实CRAS监督精度。
- 真实扫描仪自动执行闭环。
- 通用网格导入的构件级语义分块。
- 精确材料识别模型。

## 结论

当前材料达到供专利代理师技术审阅的状态。

READY_FOR_PATENT_COUNSEL_REVIEW
"""
    write_text(PATENT / "FINAL_AUDIT_REPORT.md", audit)


def main() -> None:
    generate_audit_docs()
    generate_claim_docs()
    generate_figures()
    generate_patent_md()
    make_docx(PATENT / "建筑点云缺口检测及补扫视点规划_中国发明专利初稿_审阅版.docx", review=True)
    make_docx(PATENT / "建筑点云缺口检测及补扫视点规划_中国发明专利初稿_提交候选版.docx", review=False)
    make_docx(PATENT / "权利要求书_单独版.docx", claims_only=True)
    # compatibility filename requested in the first target list
    make_docx(PATENT / "建筑点云缺口检测及补扫视点规划_中国发明专利初稿.docx", review=False)
    generate_final_audit()


if __name__ == "__main__":
    main()

#!/usr/bin/env python3
"""重新生成《权利要求书_单独版.docx》。

原文件是早期草稿脚本（generate_patent_package.py 中的 CLAIMS 常量，使用
"Patch"/"语义表面分块"等旧术语）的产物，与"最终整合版"实际提交的16项权利
要求完全不同版本，并非其同步导出。本脚本改为直接从
apply_feedback_revision.py 的 NEW_CLAIMS（反馈修订后的17项权利要求，与
最终整合版_反馈修订.docx 一致）生成单独版，格式与主文档一致（SimSun小四）。
"""

from __future__ import annotations

from pathlib import Path

import docx
from docx.shared import Pt

from apply_feedback_revision import NEW_CLAIMS

OUT = Path("/home/xqin5/patent_gap_nbv/patent_work/权利要求书_单独版.docx")


def add_body_paragraph(doc: docx.Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.5
    run = p.add_run(text)
    run.font.name = "SimSun"
    run.font.size = Pt(12)
    run.font.bold = False
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}rFonts")
    if rfonts is not None:
        rfonts.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia", "SimSun")


def add_heading_paragraph(doc: docx.Document, text: str, size: int = 14) -> None:
    p = doc.add_paragraph()
    p.alignment = 1  # center
    run = p.add_run(text)
    run.font.name = "SimSun"
    run.font.size = Pt(size)
    run.font.bold = True
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}rFonts")
    if rfonts is not None:
        rfonts.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia", "SimSun")


def main() -> None:
    doc = docx.Document()
    add_heading_paragraph(doc, "一种基于IFC表面分块和多源证据融合的建筑点云缺口检测及补充扫描视点规划方法", 14)
    add_heading_paragraph(doc, "权利要求书", 14)
    for claim in NEW_CLAIMS:
        add_body_paragraph(doc, claim)
    doc.save(str(OUT))
    print(f"saved -> {OUT} ({len(NEW_CLAIMS)} claims)")


if __name__ == "__main__":
    main()
